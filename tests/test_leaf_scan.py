from __future__ import annotations

import unicodedata
from pathlib import Path

import pytest
from pptx import Presentation

from video_workflow.leaf_scan import inspect_leaf, scan_leaf_folders


def _deck(path: Path, slides: int) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    while len(presentation.slides) < slides:
        presentation.slides.add_slide(presentation.slide_layouts[6])
    while len(presentation.slides) > slides:
        presentation.slides._sldIdLst.remove(
            presentation.slides._sldIdLst[-1]
        )
    presentation.save(path)
    return path


def _script(path: Path, text: str = "Slide 1: A") -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path


def test_valid_leaf_with_different_names(tmp_path: Path) -> None:
    leaf = tmp_path / "a" / "t40"
    _deck(leaf / "deck.pptx", 3)
    _script(
        leaf / "script.txt",
        "Slide 1: A\nSlide 2: B\nSlide 3: C",
    )

    result = scan_leaf_folders(tmp_path)

    assert len(result) == 1
    assert result[0].relpath == "a/t40"
    assert result[0].valid
    assert result[0].slide_count == 3
    assert result[0].script_chars == 3


def test_script_starting_at_slide_two_is_valid(tmp_path: Path) -> None:
    leaf = tmp_path / "lesson"
    _deck(leaf / "deck.pptx", 3)
    _script(leaf / "script.txt", "TIÊU ĐỀ\nSlide 2: B\nSlide 3: C")

    result = inspect_leaf(leaf, "lesson")

    assert result.valid
    assert result.script_chars == 2


def test_intermediate_folders_hidden(tmp_path: Path) -> None:
    parent = tmp_path / "a"
    _deck(parent / "parent.pptx", 1)
    _script(parent / "parent.txt")
    _deck(parent / "t40" / "deck.pptx", 1)
    _script(parent / "t40" / "script.txt")

    result = scan_leaf_folders(tmp_path)

    assert [leaf.relpath for leaf in result] == ["a/t40"]


def test_output_and_hidden_dirs_ignored(tmp_path: Path) -> None:
    leaf = tmp_path / "lesson"
    _deck(leaf / "deck.pptx", 1)
    _script(leaf / "script.txt")
    (leaf / "output").mkdir()
    (leaf / "output" / "x.mp4").write_bytes(b"video")
    (leaf / ".cache").mkdir()
    (leaf / ".cache" / "ignored").write_bytes(b"cache")

    result = scan_leaf_folders(tmp_path)

    assert len(result) == 1
    assert result[0].valid
    assert result[0].has_output is True


def test_junk_files_ignored(tmp_path: Path) -> None:
    leaf = tmp_path / "lesson"
    _deck(leaf / "deck.pptx", 1)
    _script(leaf / "script.txt")
    for name in ("Thumbs.db", "desktop.ini", "~$deck.pptx", ".DS_Store"):
        (leaf / name).write_bytes(b"junk")

    result = scan_leaf_folders(tmp_path)

    assert len(result) == 1
    assert result[0].valid


@pytest.mark.parametrize(
    ("files", "expected"),
    [
        (
            ("deck.pptx", "script.txt", "notes.docx"),
            "Có 3 file (cần đúng 1 .pptx và 1 .txt)",
        ),
        (("a.pptx", "b.pptx"), "Có 2 file .pptx"),
        (("a.txt", "b.txt"), "Có 2 file .txt"),
        (("script.txt",), "Thiếu file .pptx"),
        (("deck.pptx",), "Thiếu file kịch bản .txt"),
    ],
)
def test_error_messages(
    tmp_path: Path, files: tuple[str, ...], expected: str
) -> None:
    leaf = tmp_path / "lesson"
    leaf.mkdir()
    for name in files:
        path = leaf / name
        if path.suffix.casefold() == ".pptx":
            _deck(path, 1)
        elif path.suffix.casefold() == ".txt":
            _script(path)
        else:
            path.write_bytes(b"extra")

    assert inspect_leaf(leaf, "lesson").error == expected


def test_unreadable_pptx_error(tmp_path: Path) -> None:
    leaf = tmp_path / "lesson"
    leaf.mkdir()
    (leaf / "deck.pptx").write_bytes(b"not a presentation")
    _script(leaf / "script.txt")

    assert inspect_leaf(leaf, "lesson").error == "Không đọc được file .pptx"


def test_script_slide_count_error(tmp_path: Path) -> None:
    leaf = tmp_path / "lesson"
    _deck(leaf / "deck.pptx", 3)
    _script(leaf / "script.txt", "Slide 1: A\nSlide 2: B")

    assert inspect_leaf(leaf, "lesson").error == (
        "Kịch bản có 2 slide, PPTX có 3 slide"
    )


def test_other_script_error(tmp_path: Path) -> None:
    leaf = tmp_path / "lesson"
    _deck(leaf / "deck.pptx", 1)
    _script(leaf / "script.txt", "Slide 1:")

    error = inspect_leaf(leaf, "lesson").error

    assert error is not None
    assert error.startswith("Kịch bản không hợp lệ: ")


def test_non_utf8_script(tmp_path: Path) -> None:
    bad = tmp_path / "bad"
    _deck(bad / "deck.pptx", 1)
    bad.joinpath("script.txt").write_bytes("Xin chào".encode("utf-16"))
    good = tmp_path / "good"
    _deck(good / "deck.pptx", 1)
    _script(good / "script.txt")

    result = scan_leaf_folders(tmp_path)

    assert [leaf.error for leaf in result] == [
        "Kịch bản không đọc được (cần mã hoá UTF-8)",
        None,
    ]


def test_utf8_bom_script_is_valid(tmp_path: Path) -> None:
    leaf = tmp_path / "lesson"
    _deck(leaf / "deck.pptx", 1)
    leaf.joinpath("script.txt").write_bytes("Slide 1: A".encode("utf-8-sig"))

    assert inspect_leaf(leaf, "lesson").valid


def test_nfc_names(tmp_path: Path) -> None:
    nfd_name = unicodedata.normalize("NFD", "Toán")
    leaf = tmp_path / nfd_name
    _deck(leaf / "deck.pptx", 1)
    _script(leaf / "script.txt")

    result = scan_leaf_folders(tmp_path)

    assert result[0].relpath == "Toán"
    assert unicodedata.is_normalized("NFC", result[0].relpath)


def test_empty_leaf_skipped(tmp_path: Path) -> None:
    (tmp_path / "empty").mkdir()

    assert scan_leaf_folders(tmp_path) == ()
