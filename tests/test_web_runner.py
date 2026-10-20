from __future__ import annotations

import threading
from pathlib import Path

import pytest
from pptx import Presentation

from video_workflow.job_models import GlobalSettings
from video_workflow.narration import NarrationError
from video_workflow.pipeline import NarratedBuildReport, WorkflowCancelled
from video_workflow.voicestudio_client import VoiceStudioOutOfMemory
from video_workflow.web.gpu import GpuArbiter, GpuWaitCancelled
from video_workflow.web.runner import (
    RunnerContext,
    describe_error,
    run_web_job,
)
from video_workflow.web.state import WebJob


def _deck(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)


def _setup(tmp_path: Path):
    uploads = tmp_path / "uploads"
    leaf = uploads / "batch-1" / "root" / "T48_Hình thang cân"
    _deck(leaf / "TOAN6_B19_T48.pptx")
    (leaf / "script.txt").write_text("Slide 1: Xin chào.", encoding="utf-8")
    logo = tmp_path / "logo.png"
    outro = tmp_path / "outro.mp4"
    logo.write_bytes(b"logo")
    outro.write_bytes(b"outro")
    settings = GlobalSettings(logo=logo, outro=outro)
    job = WebJob.new(
        batch_id="batch-1",
        leaf_relpath="root/T48_Hình thang cân",
        profile_id="abc12345",
        script_chars=9,
    )
    captured: dict[str, object] = {}

    def build(request, **kwargs):
        captured["request"] = request
        with kwargs["synthesis_guard"]():
            kwargs["on_stage"]("synthesizing_audio")
            kwargs["on_progress"](1, 1)
        with kwargs["export_guard"]():
            kwargs["on_stage"]("exporting_slides")
        request.artifact_directory.mkdir(parents=True, exist_ok=True)
        output = request.artifact_directory / "deck.mp4"
        output.write_bytes(b"video")
        return NarratedBuildReport(
            output=str(output),
            input_signature={},
            inputs={},
            profile_config={},
            slide_count=1,
            slides=[
                {
                    "text": "Xin chào. Các em.",
                    "start_ms": 0,
                    "end_ms": 2000,
                    "generation": {"generation_seconds": 1.0},
                },
                {
                    "text": "Bài học.",
                    "start_ms": 2000,
                    "end_ms": 3000,
                    "generation": {"generation_seconds": 2.0},
                },
            ],
            total_narration_duration_ms=3000,
            outro_duration_ms=0,
            expected_duration_ms=3000,
            actual_duration_ms=3000,
            checks=[],
        )

    context = RunnerContext(
        uploads_dir=uploads,
        work_dir=tmp_path / "work",
        cache_dir=tmp_path / "cache",
        logs_dir=tmp_path / "logs",
        settings=lambda: settings,
        client=lambda: object(),
        arbiter=GpuArbiter(),
        export_lock=threading.Lock(),
        build=build,
    )
    return leaf, job, context, captured


def test_output_named_after_leaf_folder(tmp_path: Path) -> None:
    leaf, job, context, _captured = _setup(tmp_path)

    result = run_web_job(job, context, threading.Event(), lambda *_args: None)

    assert sorted(path.name for path in (leaf / "output").iterdir()) == [
        "T48_Hình thang cân.mp4",
        "vi.vtt",
    ]
    assert (leaf / "output" / "vi.vtt").read_text(encoding="utf-8").startswith(
        "WEBVTT\n\n1\n00:00:00.000"
    )
    assert result.video_path.name == "T48_Hình thang cân.mp4"
    assert result.tts_seconds == 3.0


def test_uses_job_profile_and_work_dir(tmp_path: Path) -> None:
    _leaf, job, context, captured = _setup(tmp_path)

    run_web_job(job, context, threading.Event(), lambda *_args: None)

    request = captured["request"]
    assert request.profile_config.profile_id == job.profile_id
    assert request.artifact_directory == context.work_dir / job.id
    assert request.profile_config.max_chunk_chars == 300
    assert request.profile_config.speed == 0.9


def test_stage_updates_order(tmp_path: Path) -> None:
    _leaf, job, context, _captured = _setup(tmp_path)
    stages: list[str] = []

    run_web_job(
        job,
        context,
        threading.Event(),
        lambda stage, _current, _total: stages.append(stage) if stage else None,
    )

    assert stages.index("waiting_gpu") < stages.index("synthesizing_audio")
    assert stages[-2:] == ["writing_subtitles", "publishing_output"]


def test_invalid_leaf_raises_with_message(tmp_path: Path) -> None:
    leaf, job, context, _captured = _setup(tmp_path)
    (leaf / "script.txt").unlink()

    with pytest.raises(ValueError, match="Thiếu file kịch bản .txt"):
        run_web_job(job, context, threading.Event(), lambda *_args: None)


@pytest.mark.parametrize(
    ("exc", "expected"),
    [
        (VoiceStudioOutOfMemory("CUDA OOM"), "VoiceStudio hết bộ nhớ GPU"),
        (NarrationError("slide 4 failed"), "Lỗi ở slide 4: failed"),
        (WorkflowCancelled("cancelled"), "Đã huỷ"),
        (GpuWaitCancelled("cancelled"), "Đã huỷ"),
        (ValueError("hỏng"), "Lỗi: hỏng"),
    ],
)
def test_describe_error_messages(exc: BaseException, expected: str) -> None:
    assert describe_error(exc) == expected


def test_describe_error_finds_oom_cause() -> None:
    exc = RuntimeError("wrapper")
    exc.__cause__ = VoiceStudioOutOfMemory("CUDA OOM")

    assert describe_error(exc) == "VoiceStudio hết bộ nhớ GPU"
