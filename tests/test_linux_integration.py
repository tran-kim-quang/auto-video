from __future__ import annotations

import shutil
import subprocess
import sys
import wave
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from video_workflow.pipeline import NarratedBuildRequest, build_narrated_video
from video_workflow.probe import probe_media
from video_workflow.voicestudio_client import GenerationMetadata, VoiceStudioConfig


pytestmark = pytest.mark.integration


class _WaveClient:
    """Deterministic TTS stand-in; the integration boundary starts at valid WAV."""

    def generate(
        self, text: str, destination: Path, config: VoiceStudioConfig
    ) -> GenerationMetadata:
        sample_rate = 48_000
        frame_count = sample_rate // 3
        destination.parent.mkdir(parents=True, exist_ok=True)
        with wave.open(str(destination), "wb") as output:
            output.setnchannels(1)
            output.setsampwidth(2)
            output.setframerate(sample_rate)
            output.writeframes(b"\0\0" * frame_count)
        return GenerationMetadata(
            audio_id=f"integration-{len(text)}",
            generation_seconds=0.0,
            audio_path=str(destination),
            seed=config.seed,
            audio_duration_seconds=frame_count / sample_rate,
            routing_headers={"X-Omnivoice-Device": "test"},
        )


def _require_linux_media_tools() -> None:
    if not sys.platform.startswith("linux"):
        pytest.skip("Linux integration test")
    missing = [
        tool
        for tool in ("soffice", "pdftocairo", "ffmpeg", "ffprobe")
        if shutil.which(tool) is None
    ]
    if missing:
        pytest.skip(f"missing integration tools: {', '.join(missing)}")


def _make_presentation(path: Path) -> None:
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)


def _make_outro(path: Path) -> None:
    subprocess.run(
        [
            "ffmpeg",
            "-hide_banner",
            "-loglevel",
            "error",
            "-f",
            "lavfi",
            "-i",
            "color=c=black:s=640x360:r=24:d=0.25",
            "-f",
            "lavfi",
            "-i",
            "anullsrc=r=48000:cl=stereo",
            "-shortest",
            "-c:v",
            "libx264",
            "-pix_fmt",
            "yuv420p",
            "-c:a",
            "aac",
            "-y",
            str(path),
        ],
        check=True,
    )


def test_real_linux_slide_and_ffmpeg_pipeline_builds_complete_tree(
    tmp_path: Path,
) -> None:
    _require_linux_media_tools()
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    pptx = inputs / "Bai hoc.pptx"
    script = inputs / "Bai hoc.txt"
    logo = inputs / "logo.png"
    outro = inputs / "outro.mp4"
    _make_presentation(pptx)
    script.write_text("Một. Slide 2: Hai.", encoding="utf-8")
    Image.new("RGBA", (96, 48), (255, 255, 255, 180)).save(logo)
    _make_outro(outro)

    artifact = tmp_path / "output" / pptx.stem
    report = build_narrated_video(
        NarratedBuildRequest(
            pptx=pptx,
            script=script,
            logo=logo,
            outro=outro,
            artifact_directory=artifact,
            profile_config=VoiceStudioConfig(),
        ),
        client=_WaveClient(),
        cache_root=tmp_path / "cache",
    )

    final = artifact / f"{pptx.stem}.mp4"
    info = probe_media(final)
    assert report.slide_count == 2
    assert [row["slide_id"] for row in report.slides] == [1, 2]
    assert (info.width, info.height) == (1280, 720)
    assert info.video_codec == "h264"
    assert info.audio_codec == "aac"
    assert {
        "report.json",
        "timeline.json",
        "combined.wav",
        "audio-manifest.json",
        "audio",
        "slide-images",
        final.name,
    }.issubset({path.name for path in artifact.iterdir()})
