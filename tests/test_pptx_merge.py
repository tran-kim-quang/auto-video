from __future__ import annotations

from pathlib import Path

import pytest

from video_workflow import pptx_merge


def _png(path: Path, color: str) -> Path:
    from PIL import Image

    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (1280, 720), color).save(path)
    return path


def test_merge_exports_both_parts_in_order_and_creates_40_slide_deck(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    first = tmp_path / "part1.pptx"
    second = tmp_path / "part2.pptx"
    first.touch()
    second.touch()
    calls: list[tuple[Path, list[int]]] = []

    monkeypatch.setattr(pptx_merge, "count_pptx_slides", lambda _: 20)

    def fake_export(source: Path, ids: list[int], target: Path, **_: object):
        calls.append((source, ids))
        offset = 0 if source == first else 20
        return {
            slide_id: _png(target / f"slide-{slide_id}.png", f"#{offset + slide_id:02x}0000")
            for slide_id in ids
        }

    monkeypatch.setattr(pptx_merge, "export_slides", fake_export)
    output = tmp_path / "output" / "T8_Slide.pptx"
    image_dir = tmp_path / "output" / "slide-images"

    images = pptx_merge.merge_pptx_as_images(first, second, output, image_dir)

    from pptx import Presentation

    deck = Presentation(output)
    assert len(deck.slides) == 40
    assert deck.slide_width / deck.slide_height == pytest.approx(16 / 9, rel=0.001)
    assert calls == [(first, list(range(1, 21))), (second, list(range(1, 21)))]
    assert images == tuple(image_dir / f"slide-{number:02d}.png" for number in range(1, 41))
    first_picture = deck.slides[0].shapes[0]
    assert (first_picture.left, first_picture.top) == (0, 0)
    assert (first_picture.width, first_picture.height) == (deck.slide_width, deck.slide_height)


def test_merge_requires_exactly_20_slides_per_part(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    first = tmp_path / "part1.pptx"
    second = tmp_path / "part2.pptx"
    first.touch()
    second.touch()
    monkeypatch.setattr(
        pptx_merge,
        "count_pptx_slides",
        lambda path: 19 if path == first else 20,
    )

    with pytest.raises(pptx_merge.PptxMergeError, match="part 1.*20.*19"):
        pptx_merge.merge_pptx_as_images(
            first, second, tmp_path / "merged.pptx", tmp_path / "images"
        )


def test_existing_corrupt_output_is_rejected(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    first = tmp_path / "part1.pptx"
    second = tmp_path / "part2.pptx"
    first.touch()
    second.touch()
    output = tmp_path / "merged.pptx"
    output.write_bytes(b"not a presentation")
    monkeypatch.setattr(pptx_merge, "count_pptx_slides", lambda _: 20)

    with pytest.raises(pptx_merge.PptxMergeError, match="existing merged PPTX is invalid"):
        pptx_merge.merge_pptx_as_images(first, second, output, tmp_path / "images")
