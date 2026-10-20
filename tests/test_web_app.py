from __future__ import annotations

import io
import zipfile
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from video_workflow.job_models import GlobalSettings, JobStatus
from video_workflow.voicestudio_client import VoiceStudioConfig
from video_workflow.web.app import WebServices, create_app
from video_workflow.web.config import WebLayout
from video_workflow.web.gpu import GpuArbiter
from video_workflow.web.scheduler import DuplicateJobError
from video_workflow.web.state import UploadRunRecord, WebJob, WebStore
from video_workflow.web.upload_identity import UploadCheck, upload_triple
from video_workflow.web.upload_service import UploadBusyError
from video_workflow.web.voices import VoiceInUseError, VoiceService


def _deck_bytes() -> bytes:
    buffer = io.BytesIO()
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(buffer)
    return buffer.getvalue()


class _Scheduler:
    def __init__(self) -> None:
        self.items: list[WebJob] = []

    def jobs(self):
        return tuple(self.items)

    def enqueue(self, batch_id, leaves, profile_id, output_root=None):
        for relpath, _chars in leaves:
            if any(
                job.batch_id == batch_id
                and job.leaf_relpath == relpath
                and job.status in {JobStatus.WAITING, JobStatus.RUNNING}
                for job in self.items
            ):
                raise DuplicateJobError(f"Bài {relpath} đang chờ hoặc đang chạy")
        created = [
            WebJob.new(
                batch_id=batch_id,
                leaf_relpath=relpath,
                profile_id=profile_id,
                script_chars=chars,
                output_root=output_root,
            )
            for relpath, chars in leaves
        ]
        self.items.extend(created)
        return created

    def cancel(self, job_id):
        return next(job for job in self.items if job.id == job_id)

    def retry(self, job_id):
        return next(job for job in self.items if job.id == job_id)


class _Voices:
    def list_voices(self):
        return []

    def sample_audio(self, profile_id):
        return b"audio", "audio/wav"

    def create(self, **kwargs):
        return {"id": "new"}

    def preview(self, profile_id, text, destination):
        Path(destination).write_bytes(b"preview")
        return Path(destination)

    def set_default(self, profile_id):
        return None

    def delete(self, profile_id):
        if profile_id == "busy":
            raise VoiceInUseError("Giọng đang được dùng bởi 2 bài trong hàng đợi")


class _Uploads:
    def __init__(self) -> None:
        self.busy = False
        self.items: list[UploadRunRecord] = []

    def checks(self, relpaths):
        result = {}
        for relpath in relpaths:
            triple = upload_triple(relpath)
            result[relpath] = UploadCheck(
                relpath=relpath,
                identity=None,
                issue=(
                    None
                    if triple is not None
                    else "Thiếu tầng Mã môn/Bài: cần dạng "
                    "<Mã môn>/<Bài N_Tên bài>/<T<tiết>_Tên tiết>"
                ),
                warning=None,
            )
        return result

    def start(self, batch_id, *, dry_run):
        if self.busy:
            raise UploadBusyError("Đang có một lượt upload khác chạy")
        record = UploadRunRecord(
            id="run-1",
            batch_id=batch_id,
            dry_run=dry_run,
            status="running",
            started_at="2026-10-20T10:00:00+00:00",
            finished_at=None,
            phase="video",
            current=0,
            total=1,
            video_total=1,
            video_ok=0,
            cc_total=0,
            cc_ok=0,
            video_manifest=None,
            cc_manifest=None,
            log_path=None,
            errors=[],
        )
        self.items.append(record)
        return record

    def runs(self, batch_id=None):
        return [
            item for item in self.items if batch_id is None or item.batch_id == batch_id
        ]


def _client(tmp_path: Path, voices=None):
    layout = WebLayout(tmp_path / "data").ensure()
    store = WebStore(layout.state)
    scheduler = _Scheduler()
    services = WebServices(
        layout=layout,
        store=store,
        scheduler=scheduler,
        voices=voices or _Voices(),
        uploads=_Uploads(),
        settings=GlobalSettings,
        status=lambda: {
            "voicestudio": True,
            "ffmpeg": True,
            "slide_export": True,
            "logo": "logo.png",
            "outro": "outro.mp4",
            "ready": True,
            "lan_urls": [],
            "max_reference_seconds": 7.0,
            "upload_ready": True,
            "upload_reason": None,
            "last_output_root": store.last_output_root(),
        },
    )
    return TestClient(create_app(services)), layout, scheduler


def _upload(client: TestClient):
    response = client.post(
        "/api/batches",
        files=[
            ("files", ("a.pptx", _deck_bytes(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")),
            ("files", ("a.txt", b"Slide 1: A", "text/plain")),
            ("paths", (None, "root/t40/a.pptx")),
            ("paths", (None, "root/t40/a.txt")),
        ],
    )
    assert response.status_code == 200
    return response.json()


def test_index_has_three_tabs(tmp_path: Path) -> None:
    client, _layout, _scheduler = _client(tmp_path)

    response = client.get("/")

    assert response.status_code == 200
    assert all(text in response.text for text in ("Tạo video", "Giọng đọc", "Hàng đợi"))


def test_upload_then_scan(tmp_path: Path) -> None:
    client, _layout, _scheduler = _client(tmp_path)

    batch = _upload(client)
    response = client.get(f"/api/batches/{batch['id']}")

    assert response.status_code == 200
    assert response.json()["leaves"][0]["relpath"] == "root/t40"
    assert response.json()["leaves"][0]["valid"] is True


def test_batch_leaf_includes_upload_issue_for_short_path(tmp_path: Path) -> None:
    client, _layout, _scheduler = _client(tmp_path)
    batch = _upload(client)

    leaf = client.get(f"/api/batches/{batch['id']}").json()["leaves"][0]

    assert leaf["upload_issue"].startswith("Thiếu tầng Mã môn/Bài")
    assert leaf["upload_warning"] is None
    assert leaf["upload_triple"] is None


def test_upload_busy_returns_409(tmp_path: Path) -> None:
    client, _layout, _scheduler = _client(tmp_path)
    batch = _upload(client)
    client.app.state.services.uploads.busy = True

    response = client.post(f"/api/batches/{batch['id']}/upload", json={"dry_run": True})

    assert response.status_code == 409
    assert response.json()["detail"] == "Đang có một lượt upload khác chạy"


@pytest.mark.parametrize(
    "payload",
    [{}, {"dry_run": None}, {"dry_run": "false"}, {"dry_run": 0}],
)
def test_upload_requires_explicit_strict_boolean(
    tmp_path: Path,
    payload: dict,
) -> None:
    client, _layout, _scheduler = _client(tmp_path)
    batch = _upload(client)

    response = client.post(f"/api/batches/{batch['id']}/upload", json=payload)

    assert response.status_code == 422
    assert client.app.state.services.uploads.items == []


def test_batch_leaf_upload_output_requires_exact_mp4_and_vtt(tmp_path: Path) -> None:
    client, layout, _scheduler = _client(tmp_path)
    batch = _upload(client)
    leaf = layout.uploads / batch["id"] / "root" / "t40"
    output = leaf / "output"
    output.mkdir()
    (output / "video.mp4").write_bytes(b"video")

    incomplete = client.get(f"/api/batches/{batch['id']}").json()["leaves"][0]
    assert incomplete["upload_output_ready"] is False

    (output / "vi.vtt").write_text("WEBVTT\n", encoding="utf-8")
    ready = client.get(f"/api/batches/{batch['id']}").json()["leaves"][0]
    assert ready["upload_output_ready"] is True


def test_upload_traversal_rejected(tmp_path: Path) -> None:
    client, _layout, _scheduler = _client(tmp_path)

    response = client.post(
        "/api/batches",
        files=[
            ("files", ("x.txt", b"x", "text/plain")),
            ("paths", (None, "../x.txt")),
        ],
    )

    assert response.status_code == 400
    assert response.json()["detail"]


def test_create_jobs_duplicate_409(tmp_path: Path) -> None:
    client, _layout, _scheduler = _client(tmp_path)
    batch = _upload(client)
    body = {
        "leaves": ["root/t40"],
        "profile_id": "voice",
        "output_root": str(tmp_path / "results"),
    }

    assert client.post(f"/api/batches/{batch['id']}/jobs", json=body).status_code == 200
    response = client.post(f"/api/batches/{batch['id']}/jobs", json=body)

    assert response.status_code == 409


def test_create_jobs_requires_output_root(tmp_path: Path) -> None:
    client, _layout, _scheduler = _client(tmp_path)
    batch = _upload(client)

    response = client.post(
        f"/api/batches/{batch['id']}/jobs",
        json={"leaves": ["root/t40"], "profile_id": "voice"},
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "Cần nhập thư mục lưu kết quả"


def test_create_jobs_rejects_internal_output_root(tmp_path: Path) -> None:
    client, layout, _scheduler = _client(tmp_path)
    batch = _upload(client)

    response = client.post(
        f"/api/batches/{batch['id']}/jobs",
        json={
            "leaves": ["root/t40"],
            "profile_id": "voice",
            "output_root": str(layout.uploads / "x"),
        },
    )

    assert response.status_code == 400
    assert response.json()["detail"] == (
        "Không được lưu vào thư mục nội bộ của tool"
    )


def test_create_jobs_stores_resolved_output_root_in_job_and_status(
    tmp_path: Path,
) -> None:
    client, _layout, scheduler = _client(tmp_path)
    batch = _upload(client)
    output_root = tmp_path / "Video bài giảng"

    response = client.post(
        f"/api/batches/{batch['id']}/jobs",
        json={
            "leaves": ["root/t40"],
            "profile_id": "voice",
            "output_root": str(output_root),
        },
    )

    assert response.status_code == 200
    assert scheduler.items[0].output_root == str(output_root.resolve())
    assert client.get("/api/status").json()["last_output_root"] == str(
        output_root.resolve()
    )


def test_subtitles_uses_stored_vtt_path(tmp_path: Path) -> None:
    client, _layout, scheduler = _client(tmp_path)
    video = tmp_path / "elsewhere" / "lesson.mp4"
    vtt = tmp_path / "captions" / "lesson" / "vi.vtt"
    video.parent.mkdir()
    vtt.parent.mkdir(parents=True)
    video.write_bytes(b"video")
    vtt.write_text("WEBVTT\nnew location\n", encoding="utf-8")
    job = replace(
        WebJob.new(
            batch_id="batch",
            leaf_relpath="root/t40",
            profile_id="v",
            script_chars=1,
        ),
        status=JobStatus.COMPLETED,
        video_path=str(video),
        vtt_path=str(vtt),
    )
    scheduler.items.append(job)

    response = client.get(f"/api/jobs/{job.id}/subtitles")

    assert response.status_code == 200
    assert response.text == "WEBVTT\nnew location\n"


def test_video_range_request(tmp_path: Path) -> None:
    client, _layout, scheduler = _client(tmp_path)
    video = tmp_path / "video.mp4"
    video.write_bytes(b"0123456789abcdefghij")
    job = replace(
        WebJob.new(batch_id="b", leaf_relpath="root/t40", profile_id="v", script_chars=1),
        status=JobStatus.COMPLETED,
        video_path=str(video),
    )
    scheduler.items.append(job)

    response = client.get(f"/api/jobs/{job.id}/video", headers={"Range": "bytes=0-9"})

    assert response.status_code == 206
    assert response.content == b"0123456789"


def test_zip_contains_only_completed_mp4(tmp_path: Path) -> None:
    client, _layout, scheduler = _client(tmp_path)
    video = tmp_path / "a.mp4"
    video.write_bytes(b"video")
    complete = replace(
        WebJob.new(batch_id="batch", leaf_relpath="root/t40", profile_id="v", script_chars=1),
        status=JobStatus.COMPLETED,
        video_path=str(video),
    )
    failed = replace(
        WebJob.new(batch_id="batch", leaf_relpath="root/t41", profile_id="v", script_chars=1),
        status=JobStatus.FAILED,
    )
    scheduler.items.extend([complete, failed])
    from video_workflow.web.state import BatchRecord
    scheduler_batch_store = client.app.state.services.store
    scheduler_batch_store.add_batch(BatchRecord("batch", "root", "now"))

    response = client.get("/api/batches/batch/videos.zip")

    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert archive.namelist() == ["root/t40/a.mp4"]


def test_zip_stream_does_not_buffer_archive_in_bytesio(
    tmp_path: Path,
    monkeypatch,
) -> None:
    client, _layout, scheduler = _client(tmp_path)
    video = tmp_path / "a.mp4"
    video.write_bytes(b"video")
    scheduler.items.append(
        replace(
            WebJob.new(
                batch_id="batch",
                leaf_relpath="root/t40",
                profile_id="v",
                script_chars=1,
            ),
            status=JobStatus.COMPLETED,
            video_path=str(video),
        )
    )
    from video_workflow.web.state import BatchRecord
    client.app.state.services.store.add_batch(BatchRecord("batch", "root", "now"))

    def reject_buffer(*_args, **_kwargs):
        raise AssertionError("ZIP response must not buffer the full archive")

    monkeypatch.setattr(
        "video_workflow.web.app.io",
        SimpleNamespace(BytesIO=reject_buffer),
        raising=False,
    )

    response = client.get("/api/batches/batch/videos.zip")

    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert archive.namelist() == ["root/t40/a.mp4"]


def test_delete_voice_in_use_409(tmp_path: Path) -> None:
    client, _layout, _scheduler = _client(tmp_path)

    response = client.delete("/api/voices/busy")

    assert response.status_code == 409
    assert "2 bài" in response.json()["detail"]


def test_api_voice_too_long_returns_400(tmp_path: Path) -> None:
    class ProfileClient:
        def __init__(self) -> None:
            self.created = []

        def create_profile(self, **kwargs):
            self.created.append(kwargs)
            return {"id": "new", "name": kwargs["name"], "kind": "clone"}

    profile_client = ProfileClient()
    store = WebStore(tmp_path / "voice-state")
    voices = VoiceService(
        client=lambda: profile_client,
        arbiter=GpuArbiter(),
        store=store,
        config=VoiceStudioConfig,
        active_jobs=tuple,
        max_reference_seconds=7,
        probe=lambda _path: SimpleNamespace(duration_ms=10000),
    )
    client, _layout, _scheduler = _client(tmp_path, voices=voices)

    response = client.post(
        "/api/voices",
        data={"name": "Cô Lan", "ref_text": "xin chào"},
        files={"audio": ("lan.wav", b"audio", "audio/wav")},
    )

    assert response.status_code == 400
    assert response.json()["detail"] == (
        "File ghi âm dài 10.0 giây, tối đa 7 giây trên máy này — "
        "hãy cắt ngắn và nhập đúng lời nói của đoạn đó"
    )
    assert profile_client.created == []


def test_no_vtt_download_link_in_ui() -> None:
    static = Path(__file__).parents[1] / "video_workflow" / "web" / "static" / "app.js"
    assert "Tải vi.vtt" not in static.read_text(encoding="utf-8")


def test_static_assets_served(tmp_path: Path) -> None:
    client, _layout, _scheduler = _client(tmp_path)

    assert client.get("/static/app.js").status_code == 200
    assert client.get("/static/styles.css").status_code == 200


def test_stage_labels_present() -> None:
    source = (
        Path(__file__).parents[1]
        / "video_workflow"
        / "web"
        / "static"
        / "app.js"
    ).read_text(encoding="utf-8")

    for label in (
        "Đang chờ",
        "Chờ GPU",
        "Chờ xuất slide",
        "Kiểm tra đầu vào",
        "Sinh giọng",
        "Ghép audio",
        "Xuất slide",
        "Render video",
        "Chuẩn bị outro",
        "Ghép outro",
        "Kiểm tra video",
        "Đóng gói",
        "Tạo phụ đề",
        "Lưu kết quả",
    ):
        assert label in source


def test_default_preview_text_present() -> None:
    source = (
        Path(__file__).parents[1]
        / "video_workflow"
        / "web"
        / "static"
        / "index.html"
    ).read_text(encoding="utf-8")

    assert "Cho tam giác ABC vuông tại A, tính độ dài cạnh BC." in source


def test_ui_hint_uses_status_limit() -> None:
    static_dir = Path(__file__).parents[1] / "video_workflow" / "web" / "static"
    source = "\n".join(
        path.read_text(encoding="utf-8")
        for path in (static_dir / "index.html", static_dir / "app.js")
    )

    assert "Nên 5–15 giây" not in source
    assert "max_reference_seconds" in source


def test_r2_upload_controls_and_results_are_present() -> None:
    source = (
        Path(__file__).parents[1]
        / "video_workflow"
        / "web"
        / "static"
        / "app.js"
    ).read_text(encoding="utf-8")

    for text in (
        "Kiểm tra upload",
        "Upload lên R2",
        "KHÔNG import",
        "Manifest phụ đề",
    ):
        assert text in source


def test_r2_controls_use_exact_output_eligibility_and_safe_copy_fallback() -> None:
    source = (
        Path(__file__).parents[1]
        / "video_workflow"
        / "web"
        / "static"
        / "app.js"
    ).read_text(encoding="utf-8")

    assert "upload_output_ready" in source
    assert "l?.has_output" not in source
    assert "new Set" in source
    assert 'document.execCommand("copy")' in source
    assert "upload_warning" in source


def test_output_root_controls_and_saved_path_are_present() -> None:
    static_dir = Path(__file__).parents[1] / "video_workflow" / "web" / "static"
    index = (static_dir / "index.html").read_text(encoding="utf-8")
    javascript = (static_dir / "app.js").read_text(encoding="utf-8")

    assert "Thư mục lưu kết quả (trên máy server)" in index
    assert 'id="output-root"' in index
    assert "output_root" in javascript
    assert "Lưu tại:" in javascript
