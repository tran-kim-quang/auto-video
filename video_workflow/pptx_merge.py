from __future__ import annotations

import shutil
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from .slide_export import count_pptx_slides, export_slides


class PptxMergeError(RuntimeError):
    """Raised when the two lesson decks cannot form one verified deck."""


def _open_40_slide_deck(path: Path) -> Presentation:
    try:
        deck = Presentation(path)
    except Exception as exc:
        raise PptxMergeError(f"existing merged PPTX is invalid: {path}") from exc
    if len(deck.slides) != 40:
        raise PptxMergeError(
            f"existing merged PPTX is invalid: expected 40 slides, found {len(deck.slides)}"
        )
    return deck


def merge_pptx_as_images(
    first: Path,
    second: Path,
    output: Path,
    image_dir: Path,
) -> tuple[Path, ...]:
    first = Path(first)
    second = Path(second)
    output = Path(output)
    image_dir = Path(image_dir)
    references = tuple(image_dir / f"slide-{number:02d}.png" for number in range(1, 41))

    for label, source in (("part 1", first), ("part 2", second)):
        count = count_pptx_slides(source)
        if count != 20:
            raise PptxMergeError(f"{label} must contain 20 slides, found {count}")

    if output.exists():
        _open_40_slide_deck(output)
        if all(path.is_file() and path.stat().st_size > 0 for path in references):
            return references
        exported = export_slides(output, list(range(1, 41)), image_dir, width=1280, height=720)
        for slide_id, destination in enumerate(references, start=1):
            source = exported[slide_id]
            if source != destination:
                shutil.copy2(source, destination)
        return references

    output.parent.mkdir(parents=True, exist_ok=True)
    image_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="pptx-merge-", dir=output.parent) as temporary:
        staging = Path(temporary)
        first_images = export_slides(
            first, list(range(1, 21)), staging / "part1", width=1280, height=720
        )
        second_images = export_slides(
            second, list(range(1, 21)), staging / "part2", width=1280, height=720
        )
        ordered = [first_images[number] for number in range(1, 21)]
        ordered += [second_images[number] for number in range(1, 21)]

        deck = Presentation()
        deck.slide_width = Inches(13.333333)
        deck.slide_height = Inches(7.5)
        blank_layout = deck.slide_layouts[6]
        for source in ordered:
            slide = deck.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                str(source), 0, 0, width=deck.slide_width, height=deck.slide_height
            )

        temporary_pptx = output.parent / f".{output.name}.tmp"
        try:
            deck.save(temporary_pptx)
            _open_40_slide_deck(temporary_pptx)
            temporary_pptx.replace(output)
        finally:
            temporary_pptx.unlink(missing_ok=True)

        for number, source in enumerate(ordered, start=1):
            shutil.copy2(source, references[number - 1])
    return references
